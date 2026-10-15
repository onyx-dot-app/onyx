"""Production document tools and frontend delivery on a normally provisioned guest."""

from __future__ import annotations

import hashlib
import io
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable
from pathlib import Path, PurePosixPath
from typing import TYPE_CHECKING
from uuid import UUID, uuid4

from docx import Document
from PIL import Image, ImageStat
from pydantic import BaseModel, ConfigDict
from pypdf import PdfReader

from onyx.server.features.build.configs import SANDBOX_NAMESPACE
from onyx.server.features.build.sandbox.manager import SandboxManager
from onyx.server.features.build.session.models import PptxPreviewResponse
from tests.integration.common_utils.constants import API_SERVER_URL
from tests.integration.common_utils.http_client import client
from tests.integration.common_utils.managers.build_session import BuildSessionManager
from tests.integration.tests.craft.k8s.k8s_fixtures import (
    OwnedLivePod,
    kubernetes_runtime,
)

if TYPE_CHECKING:
    from kubernetes.client import CoreV1Api


class _Artifact(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)

    size_bytes: int
    sha256: str


class _ProtectedPath(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)

    path: str
    uid: int
    writable: bool


class _UserSite(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)

    interpreter: str
    uid: int
    euid: int
    gid: int
    home: str
    path: str
    enabled: bool
    in_sys_path: bool
    module_path: str
    module_uid: int
    value: str


class _Receipt(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)

    uid: int
    euid: int
    gid: int
    home: str
    home_uid: int
    interpreter: str
    boot_id: str
    user_site: _UserSite
    protected_paths: list[_ProtectedPath]
    files: dict[str, _Artifact]


_PRESENTATION_SCRIPT = r"""
"use strict";
const path = require("path");
const C = require(path.resolve(process.cwd(), ".opencode/skills/pptx/scripts/components.js"));
const nonce = process.argv[2];
const theme = C.makeTheme({
  primary: "1E2761", darkField: "1E2761", gridDark: "1E2761", background: "FFFFFF"
});
const pres = C.createPresentation(theme, { title: `Craft ${nonce}` });
C.sectionDivider(pres, theme, { number: 1, title: `Deck ${nonce}` });
C.iconRow(pres, theme, {
  title: `Tools ${nonce}`,
  items: [
    { iconPath: `outputs/icon-${nonce}.png`, label: "Production icon" },
    { label: "Document tools", desc: `Craft ${nonce}` }
  ]
});
pres.writeFile({ fileName: `outputs/deck-${nonce}.pptx` }).catch(error => {
  console.error(error);
  process.exitCode = 1;
});
"""

_DOCUMENT_SCRIPT = r"""
import hashlib
import json
import os
import subprocess
import sys
from pathlib import Path

nonce, presentation_script = sys.argv[1:]
expected = f"Craft {nonce}"
outputs = Path("outputs")
outputs.mkdir(exist_ok=True)
skill_scripts = Path(".opencode/skills/pptx/scripts")
for name in ("components.js", "icon.js", "chart.py"):
    if not (skill_scripts / name).is_file():
        raise RuntimeError(f"Provisioning did not push required PPTX script: {name}")

user_site = Path(subprocess.check_output(
    ["/usr/local/bin/python3", "-c", "import site; print(site.getusersitepackages())"],
    text=True, timeout=20,
).strip())
assert user_site.is_relative_to(Path("/home/sandbox/.local/lib")), user_site
user_site.mkdir(parents=True, exist_ok=True)
module_name = f"craft_production_{nonce}"
module_path = user_site / f"{module_name}.py"
module_path.write_text(f"VALUE = {expected!r}\n", encoding="utf-8")
site_result = json.loads(subprocess.check_output(
    ["/usr/local/bin/python3", "-c", '''
import importlib
import importlib.util
import json
import os
import site
import sys
from pathlib import Path
spec = importlib.util.find_spec(sys.argv[1])
assert spec is not None and spec.origin is not None, "User-site module is not importable"
module = importlib.import_module(sys.argv[1])
print(json.dumps({
    "interpreter": sys.executable, "uid": os.getuid(), "euid": os.geteuid(),
    "gid": os.getgid(), "home": str(Path.home()), "path": site.getusersitepackages(),
    "enabled": site.ENABLE_USER_SITE, "in_sys_path": site.getusersitepackages() in sys.path,
    "module_path": spec.origin, "module_uid": Path(spec.origin).stat().st_uid,
    "value": module.VALUE,
}))
''', module_name], text=True, timeout=20,
))
protected_paths = json.loads(subprocess.check_output(
    ["/opt/onyx-guest/venv/bin/python", "-I", "-c", '''
import importlib.util
import json
import os
from pathlib import Path
spec = importlib.util.find_spec("onyx_guest")
assert spec is not None and spec.origin is not None
roots = [Path("/opt/onyx-guest/venv/bin/python"), Path(spec.origin)]
paths = {path for root in roots for path in (root, root.resolve(), *root.parents, *root.resolve().parents)}
print(json.dumps([
    {"path": str(path), "uid": path.stat().st_uid, "writable": os.access(path, os.W_OK)}
    for path in sorted(paths)
]))
'''], text=True, timeout=20,
))
for protected in protected_paths:
    assert protected["uid"] == 0 and not protected["writable"], protected

try:
    for icon_set, icon_name, prefix in (
        ("lucide", "circle-check", "icon"),
        ("tabler", "shield-check", "tabler"),
    ):
        subprocess.run(
            ["node", str(skill_scripts / "icon.js"), icon_name, "--set", icon_set,
             "--color", "1E2761", "--size", "128", "-o", f"outputs/{prefix}-{nonce}.png"],
            check=True, capture_output=True, text=True, timeout=30,
        )
    subprocess.run(
        ["node", presentation_script, nonce],
        check=True, capture_output=True, text=True, timeout=30,
    )
except subprocess.CalledProcessError as error:
    raise RuntimeError(f"Production PPTX tooling failed: {error.stderr}") from error

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from matplotlib.figure import Figure
from openpyxl import Workbook

sys.path.insert(0, str(skill_scripts.resolve()))
from chart import deck_style, save_for_slide

with deck_style(palette=["1E2761"], font="DejaVu Sans", background="FFFFFF"):
    fig, ax = plt.subplots(layout="constrained")
    ax.bar(["Craft"], [7])
    ax.set_title(expected)
    save_for_slide(fig, f"outputs/chart-{nonce}.png", width_in=4, height_in=2)

pdf = Figure(figsize=(8, 4.5), facecolor="#1E2761")
pdf.text(0.1, 0.9, expected, color="white")
pdf.add_axes((0.1, 0.15, 0.6, 0.6)).imshow(plt.imread(f"outputs/chart-{nonce}.png"))
pdf.axes[0].set_axis_off()
pdf.savefig(f"outputs/production-{nonce}.pdf", format="pdf", facecolor="#1E2761")

workbook = Workbook()
workbook.active["A1"] = expected
workbook.active["B1"] = 7
workbook.save(f"outputs/workbook-{nonce}.xlsx")
document = Document()
document.add_paragraph(site_result["value"])
document.save(f"outputs/document-{nonce}.docx")

files = {}
for path in sorted(outputs.glob(f"*-{nonce}.*")):
    content = path.read_bytes()
    files[str(path)] = {"size_bytes": len(content), "sha256": hashlib.sha256(content).hexdigest()}
print(json.dumps({
    "uid": os.getuid(), "euid": os.geteuid(), "gid": os.getgid(),
    "home": str(Path.home()), "home_uid": Path.home().stat().st_uid,
    "interpreter": sys.executable,
    "boot_id": Path("/proc/sys/kernel/random/boot_id").read_text().strip(),
    "user_site": site_result, "protected_paths": protected_paths, "files": files,
}))
"""


def test_production_image_generates_and_previews_document_outputs(
    owned_live_pod: OwnedLivePod,
    k8s_manager: SandboxManager,
    k8s_client: CoreV1Api,
    record_property: Callable[[str, object], None],
    tmp_path: Path,
) -> None:
    pod = k8s_client.read_namespaced_pod(owned_live_pod.pod_name, SANDBOX_NAMESPACE)
    assert pod.spec.runtime_class_name == "kata-fc"
    statuses = pod.status.container_statuses
    assert statuses is not None
    image_status = next(status for status in statuses if status.name == "sandbox")
    assert image_status.image_id is not None
    assert "sha256:" in image_status.image_id, image_status.image_id
    record_property("sandbox_pod_uid", str(pod.metadata.uid))
    record_property("sandbox_image", image_status.image)
    record_property("sandbox_image_id", image_status.image_id)

    user, session_id = owned_live_pod.api_user, owned_live_pod.session_id
    nonce = uuid4().hex[:12]
    expected = f"Craft {nonce}"
    presentation = BuildSessionManager.upload_file(
        user, session_id, f"generate-{nonce}.js", _PRESENTATION_SCRIPT.encode()
    )
    generation = BuildSessionManager.upload_file(
        user, session_id, f"generate-{nonce}.py", _DOCUMENT_SCRIPT.encode()
    )
    markdown = BuildSessionManager.upload_file(
        user,
        session_id,
        f"export-{nonce}.md",
        f"# {expected}\n\nProduction export {nonce}.\n\n| Name | Count |\n| --- | --- |\n| {nonce} | 7 |\n".encode(),
    )
    result = kubernetes_runtime(k8s_manager).helper(
        owned_live_pod.sandbox_id,
        "run",
        {
            "argv": [
                "/workspace/.venv/bin/python",
                generation.path,
                nonce,
                presentation.path,
            ],
            "cwd": f"/workspace/sessions/{session_id}",
            "timeout_seconds": 180,
            "has_stdin": False,
        },
        timeout=210,
    )
    assert result.returncode == 0, result.stderr
    receipt = _Receipt.model_validate_json(result.stdout)
    record_property("production_document_receipt", receipt.model_dump_json())
    assert (receipt.uid, receipt.euid, receipt.gid, receipt.home_uid) == (
        1000,
        1000,
        1000,
        1000,
    )
    assert receipt.home == "/home/sandbox"
    assert receipt.interpreter == "/workspace/.venv/bin/python"
    UUID(receipt.boot_id)
    site = receipt.user_site
    assert (site.uid, site.euid, site.gid, site.module_uid) == (1000, 1000, 1000, 1000)
    assert site.interpreter == "/usr/local/bin/python3"
    assert site.home == "/home/sandbox"
    assert site.enabled is True and site.in_sys_path is True
    assert PurePosixPath(site.path).is_relative_to("/home/sandbox/.local/lib")
    assert site.module_path == f"{site.path}/craft_production_{nonce}.py"
    assert site.value == expected
    assert len(receipt.protected_paths) >= 4
    for protected in receipt.protected_paths:
        assert protected.uid == 0 and protected.writable is False, protected

    paths = {
        "deck": f"outputs/deck-{nonce}.pptx",
        "icon": f"outputs/icon-{nonce}.png",
        "tabler": f"outputs/tabler-{nonce}.png",
        "chart": f"outputs/chart-{nonce}.png",
        "pdf": f"outputs/production-{nonce}.pdf",
        "workbook": f"outputs/workbook-{nonce}.xlsx",
        "document": f"outputs/document-{nonce}.docx",
    }
    assert set(receipt.files) == set(paths.values())
    downloads: dict[str, bytes] = {}
    for kind, path in paths.items():
        content = BuildSessionManager.download_artifact(user, session_id, path)
        downloads[kind] = content
        (tmp_path / PurePosixPath(path).name).write_bytes(content)
        assert len(content) == receipt.files[path].size_bytes
        assert hashlib.sha256(content).hexdigest() == receipt.files[path].sha256

    with zipfile.ZipFile(io.BytesIO(downloads["deck"])) as deck:
        for number, text in ((1, f"Deck {nonce}"), (2, f"Tools {nonce}")):
            slide = ET.fromstring(deck.read(f"ppt/slides/slide{number}.xml"))
            assert text in [
                node.text for node in slide.iter() if node.tag.endswith("}t")
            ]
        media = [name for name in deck.namelist() if name.startswith("ppt/media/")]
        assert downloads["icon"] in [deck.read(name) for name in media]
    with zipfile.ZipFile(io.BytesIO(downloads["workbook"])) as workbook:
        sheet = ET.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        assert (
            sheet.findtext(".//s:c[@r='A1']/s:is/s:t", namespaces=namespace) == expected
        )
        assert sheet.findtext(".//s:c[@r='B1']/s:v", namespaces=namespace) == "7"
    assert [p.text for p in Document(io.BytesIO(downloads["document"])).paragraphs] == [
        expected
    ]
    pdf = PdfReader(io.BytesIO(downloads["pdf"]))
    assert len(pdf.pages) == 1
    assert expected in pdf.pages[0].extract_text()

    for kind, size in (
        ("icon", (128, 128)),
        ("tabler", (128, 128)),
        ("chart", (800, 400)),
    ):
        with Image.open(io.BytesIO(downloads[kind])) as image:
            assert image.format == "PNG"
            assert image.size == size
            assert max(ImageStat.Stat(image.convert("RGB")).stddev) > 10

    base_url = f"{API_SERVER_URL}/build/sessions/{session_id}"
    response = client.get(
        f"{base_url}/pptx-preview/{paths['deck']}",
        headers=user.headers,
        cookies=user.cookies,
    )
    assert response.status_code == 200, response.text
    preview = PptxPreviewResponse.model_validate(response.json())
    assert preview.slide_count == 2
    assert preview.cached is False
    cache_hash = hashlib.sha256(paths["deck"].encode()).hexdigest()[:12]
    assert len(preview.slide_paths) == 2
    for number, path in enumerate(preview.slide_paths, start=1):
        relative_path = PurePosixPath(path)
        assert relative_path.parent == PurePosixPath(
            f"outputs/.pptx-preview/{cache_hash}"
        )
        assert int(relative_path.stem.removeprefix("slide-")) == number
        content = BuildSessionManager.download_artifact(user, session_id, path)
        (tmp_path / f"preview-{number}.jpg").write_bytes(content)
        with Image.open(io.BytesIO(content)) as image:
            assert image.format == "JPEG"
            assert abs(image.width / image.height - 16 / 9) < 0.02
            pixel = image.convert("RGB").getpixel(
                (int(image.width * 0.97), int(image.height * 0.05))
            )
            color = (30, 39, 97) if number == 1 else (255, 255, 255)
            assert isinstance(pixel, tuple)
            for actual, target in zip(pixel, color, strict=True):
                assert isinstance(actual, int)
                assert abs(actual - target) <= 15, pixel

    response = client.get(
        f"{base_url}/output-thumbnail/{paths['pdf']}",
        headers=user.headers,
        cookies=user.cookies,
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "image/jpeg"
    (tmp_path / "pdf-thumbnail.jpg").write_bytes(response.content)
    with Image.open(io.BytesIO(response.content)) as image:
        assert image.format == "JPEG"
        assert max(image.size) == 640
        pixel = image.convert("RGB").getpixel(
            (int(image.width * 0.97), int(image.height * 0.05))
        )
        assert isinstance(pixel, tuple)
        for actual, target in zip(pixel, (30, 39, 97), strict=True):
            assert isinstance(actual, int)
            assert abs(actual - target) <= 15, pixel
    assert (
        BuildSessionManager.download_artifact(user, session_id, paths["pdf"])
        == downloads["pdf"]
    )

    response = client.get(
        f"{base_url}/export-docx/{markdown.path}",
        headers=user.headers,
        cookies=user.cookies,
    )
    assert response.status_code == 200, response.text
    assert (
        response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert (
        response.headers["content-disposition"]
        == f'attachment; filename="export-{nonce}.docx"; '
        f"filename*=UTF-8''export-{nonce}.docx"
    )
    (tmp_path / "export.docx").write_bytes(response.content)
    exported = Document(io.BytesIO(response.content))
    assert [paragraph.text for paragraph in exported.paragraphs if paragraph.text] == [
        expected,
        f"Production export {nonce}.",
    ]
    assert exported.paragraphs[0].style is not None
    assert exported.paragraphs[0].style.name == "Heading 1"
    assert [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in exported.tables
    ] == [[["Name", "Count"], [nonce, "7"]]]
