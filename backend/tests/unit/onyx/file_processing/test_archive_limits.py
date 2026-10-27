import asyncio
import os
import time
import tracemalloc
import zipfile
from io import BytesIO
from typing import IO, Literal, cast
from unittest.mock import MagicMock, patch
from uuid import uuid4

import openpyxl
import pytest
from fastapi import UploadFile
from openpyxl.worksheet.worksheet import Worksheet

from onyx.configs.constants import ONYX_METADATA_FILENAME
from onyx.db import projects
from onyx.db.models import User
from onyx.error_handling.error_codes import OnyxErrorCode
from onyx.error_handling.exceptions import OnyxError
from onyx.file_processing import extract_file_text as extraction
from onyx.file_processing import zip_limits
from onyx.server.documents import connector
from onyx.server.features.build.user_library import api as library
from onyx.server.query_and_chat import chat_utils
from onyx.skills import metadata


def _archive(
    entries: dict[str, bytes], compression: int = zipfile.ZIP_STORED
) -> BytesIO:
    buffer: BytesIO = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=compression) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    buffer.seek(0)
    return buffer


@pytest.mark.parametrize("limit", ["member", "total", "ratio"])
def test_archive_preflight_rejects_expansion_before_read(
    monkeypatch: pytest.MonkeyPatch, limit: str
) -> None:
    if limit == "member":
        monkeypatch.setattr(zip_limits, "MAX_ZIP_MEMBER_DECOMPRESSED_BYTES", 8)
        good: BytesIO = _archive({"file.txt": b"x" * 8})
        bad: BytesIO = _archive({"file.txt": b"x" * 9})
    elif limit == "total":
        monkeypatch.setattr(zip_limits, "MAX_ZIP_TOTAL_DECOMPRESSED_BYTES", 8)
        good = _archive({"a.txt": b"1234", "b.txt": b"5678"})
        bad = _archive({"a.txt": b"1234", "b.txt": b"56789"})
    else:
        monkeypatch.setattr(zip_limits, "_ZIP_RATIO_CHECK_MIN_BYTES", 16)
        good = _archive({"file.txt": b"x" * 4096})
        bad = _archive({"file.txt": b"x" * 4096}, zipfile.ZIP_DEFLATED)
    with zipfile.ZipFile(good) as archive:
        zip_limits.assert_zip_within_limits(archive)
    with zipfile.ZipFile(bad) as archive, patch.object(archive, "open") as open_member:
        with pytest.raises(ValueError, match="Zip"):
            zip_limits.assert_zip_within_limits(archive)
        open_member.assert_not_called()


def test_member_read_bounds_actual_bytes_even_with_inaccurate_metadata() -> None:
    archive: MagicMock = MagicMock()
    info: zipfile.ZipInfo = zipfile.ZipInfo("file.txt")
    info.file_size = 4
    stream: MagicMock = MagicMock(wraps=BytesIO(b"123456789"))
    archive.open.return_value.__enter__.return_value = stream
    with pytest.raises(ValueError, match="too large"):
        zip_limits.read_zip_member(archive, info, max_bytes=4)
    stream.read.assert_called_once_with(5)
    archive.open.return_value.__exit__.assert_called_once()


@pytest.mark.parametrize("compression", [zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA])
def test_archive_rejects_compression_without_bounded_reads(compression: int) -> None:
    with zipfile.ZipFile(_archive({"file.txt": b"x"}, compression)) as archive:
        info: zipfile.ZipInfo = archive.infolist()[0]
        with patch.object(archive, "open") as open_member:
            with pytest.raises(ValueError, match="compression"):
                zip_limits.assert_zip_within_limits(archive)
            with pytest.raises(ValueError, match="compression"):
                zip_limits.read_zip_member(archive, info)
            open_member.assert_not_called()


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_office_text_preflight_stops_converter(
    monkeypatch: pytest.MonkeyPatch, kind: str
) -> None:
    monkeypatch.setattr(zip_limits, "MAX_ZIP_MEMBER_DECOMPRESSED_BYTES", 8)
    reader = extraction.read_docx_file if kind == "docx" else extraction.pptx_to_text
    buffer: BytesIO = _archive({"document.xml": b"x" * 9})
    buffer.seek(2)
    with patch.object(extraction, "get_markitdown_converter") as converter:
        with pytest.raises(ValueError, match="too large"):
            reader(buffer)
        converter.return_value.convert.assert_not_called()
    assert buffer.tell() == 2


@pytest.mark.parametrize("kind", ["docx", "pptx"])
def test_office_images_skip_oversized_archive(
    monkeypatch: pytest.MonkeyPatch, kind: str
) -> None:
    monkeypatch.setattr(zip_limits, "MAX_ZIP_MEMBER_DECOMPRESSED_BYTES", 8)
    reader = (
        extraction.extract_docx_images
        if kind == "docx"
        else extraction.extract_pptx_images
    )
    path: str = "word/media/image.png" if kind == "docx" else "ppt/media/image.png"
    assert list(reader(_archive({path: b"12345678"}))) == [(b"12345678", "image.png")]
    assert list(reader(_archive({path: b"123456789"}))) == []


def test_epub_rejects_expansion_and_accepts_normal_html(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(zip_limits, "MAX_ZIP_MEMBER_DECOMPRESSED_BYTES", 64)
    assert "Hello" in extraction.epub_to_text(
        _archive({"chapter.html": b"<p>Hello</p>"})
    )
    with pytest.raises(ValueError, match="too large"):
        extraction.epub_to_text(
            _archive({"chapter.html": b"<p>" + b"x" * 64 + b"</p>"})
        )


def _workbook_bytes() -> bytes:
    workbook: openpyxl.Workbook = openpyxl.Workbook()
    cast(Worksheet, workbook.active).append(["a", "b"])
    buffer: BytesIO = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xlsx_bomb() -> BytesIO:
    """A real xlsx plus a 256 MiB sharedStrings.xml that deflates to ~250 KB."""
    buffer: BytesIO = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(_workbook_bytes())) as source,
        zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            target.writestr(info.filename, source.read(info))
        with target.open("xl/sharedStrings.xml", "w", force_zip64=True) as member:
            chunk: bytes = b"\0" * (1024 * 1024)
            for _ in range(256):
                member.write(chunk)
    buffer.seek(0)
    return buffer


def _recompress(data: bytes, compression: int) -> BytesIO:
    buffer: BytesIO = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(data)) as source,
        zipfile.ZipFile(buffer, "w", compression) as target,
    ):
        for info in source.infolist():
            target.writestr(info.filename, source.read(info))
    buffer.seek(0)
    return buffer


@pytest.mark.parametrize("entry", ["sheets", "stage", "text"])
def test_xlsx_bomb_rejected_before_openpyxl(entry: str) -> None:
    bomb: BytesIO = _xlsx_bomb()
    with patch.object(extraction.openpyxl, "load_workbook") as load_workbook:
        started: float = time.monotonic()
        with pytest.raises(zip_limits.ZipSizeLimitError):
            if entry == "sheets":
                extraction.xlsx_sheet_extraction(bomb, "bomb.xlsx")
            elif entry == "stage":
                extraction.stage_xlsx_sheets(bomb, MagicMock(), "bomb.xlsx")
            else:
                extraction.xlsx_to_text(bomb, "bomb.xlsx")
        assert time.monotonic() - started < 5
        load_workbook.assert_not_called()


def test_xlsx_rejects_bzip2_members() -> None:
    xlsx: BytesIO = _recompress(_workbook_bytes(), zipfile.ZIP_BZIP2)
    with patch.object(extraction.openpyxl, "load_workbook") as load_workbook:
        with pytest.raises(zip_limits.ZipSizeLimitError, match="compression"):
            extraction.xlsx_sheet_extraction(xlsx, "bzip2.xlsx")
        load_workbook.assert_not_called()


def test_xlsx_preview_uses_tighter_decompression_cap(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data: bytes = _workbook_bytes()
    assert extraction.xlsx_sheet_extraction(BytesIO(data), "ok.xlsx")[0][0] == "a,b"
    monkeypatch.setattr(
        chat_utils, "MAX_PREVIEW_DECOMPRESSED_BYTES", 1024, raising=False
    )
    with pytest.raises(OnyxError, match="too large") as error:
        chat_utils.parse_spreadsheet_for_preview(BytesIO(data), "ok.xlsx")
    assert error.value.error_code == OnyxErrorCode.INVALID_INPUT


def test_xlsx_sparse_wide_rows_stay_within_cell_budget(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(extraction, "MAX_XLSX_CELLS_PER_SHEET", 1_000_000)
    workbook: openpyxl.Workbook = openpyxl.Workbook(write_only=True)
    sheet = workbook.create_sheet()
    padding: list[None] = [None] * 16383
    for _ in range(2000):
        sheet.append(padding + ["1"])
    buffer: BytesIO = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    tracemalloc.start()
    try:
        sheets = extraction.xlsx_sheet_extraction(buffer, "wide.xlsx")
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert sheets[0][0].endswith("[truncated: sheet exceeded cell limit]")
    assert peak < 32 * 1024 * 1024


@pytest.mark.parametrize("compression", ["bzip2", "oversized"])
def test_zip_loader_rejects_unbounded_members(
    monkeypatch: pytest.MonkeyPatch, compression: str
) -> None:
    if compression == "bzip2":
        archive: BytesIO = _archive(
            {"site/PUBLISHED/index.html": b"<p>x</p>"}, zipfile.ZIP_BZIP2
        )
    else:
        monkeypatch.setattr(zip_limits, "MAX_ZIP_MEMBER_DECOMPRESSED_BYTES", 8)
        archive = _archive({"site/PUBLISHED/index.html": b"x" * 9})
    with patch.object(zipfile.ZipFile, "open") as open_member:
        with pytest.raises(zip_limits.ZipSizeLimitError):
            list(extraction.load_files_from_zip(archive))
        open_member.assert_not_called()


def test_zip_loader_yields_member_contents() -> None:
    archive: BytesIO = _archive({"a/PUBLISHED/index.html": b"<p>x</p>"})
    files = [
        (info.filename, handle.read())
        for info, handle in extraction.load_files_from_zip(archive)
    ]
    assert files == [("a/PUBLISHED/index.html", b"<p>x</p>")]


def test_zip_loader_does_not_read_skipped_members() -> None:
    archive: BytesIO = _archive(
        {"a/attachment.bin": os.urandom(4 * 1024 * 1024), "a/index.html": b"<p></p>"}
    )
    tracemalloc.start()
    try:
        names = [info.filename for info, _ in extraction.load_files_from_zip(archive)]
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert names == ["a/attachment.bin", "a/index.html"]
    assert peak < 1024 * 1024


def test_xlsx_preview_flags_cell_cap_truncation(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(extraction, "MAX_XLSX_CELLS_PER_SHEET", 4)
    workbook: openpyxl.Workbook = openpyxl.Workbook()
    sheet = cast(Worksheet, workbook.active)
    for row in range(5):
        sheet.append(["x", str(row)])
    buffer: BytesIO = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    preview = chat_utils.parse_spreadsheet_for_preview(buffer, "wide.xlsx")
    assert preview.sheets[0].truncated
    assert preview.sheets[0].csv == "x,0\nx,1\nx,2"


def test_connector_checks_limits_when_storing_zip_as_is(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(connector, "MAX_UNZIPPED_BYTES", 8)
    upload: UploadFile = UploadFile(
        file=_archive({"a.txt": b"12345", "b.txt": b"6789"}), filename="site.zip"
    )
    with patch.object(connector, "get_default_file_store") as store:
        store.return_value.save_file.return_value = "stored-file"
        with pytest.raises(OnyxError, match="too large") as error:
            connector.upload_files([upload], unzip=False)
        assert error.value.error_code == OnyxErrorCode.INVALID_INPUT
        store.return_value.save_file.assert_not_called()


@pytest.mark.parametrize("metadata_member", [False, True])
def test_connector_rejects_total_expansion_before_storage(
    monkeypatch: pytest.MonkeyPatch, metadata_member: bool
) -> None:
    monkeypatch.setattr(connector, "MAX_UNZIPPED_BYTES", 8)
    name: str = ONYX_METADATA_FILENAME if metadata_member else "a.txt"
    upload: UploadFile = UploadFile(
        file=_archive({name: b"12345", "b.txt": b"6789"}), filename="files.zip"
    )
    with patch.object(connector, "get_default_file_store") as store:
        store.return_value.save_file.return_value = "stored-file"
        with pytest.raises(OnyxError, match="too large") as error:
            connector.upload_files([upload])
        assert error.value.error_code == OnyxErrorCode.INVALID_INPUT
        store.return_value.save_file.assert_not_called()


def test_connector_rejects_invalid_metadata_json() -> None:
    upload: UploadFile = UploadFile(
        file=_archive({ONYX_METADATA_FILENAME: b"{"}), filename="files.zip"
    )
    with patch.object(connector, "get_default_file_store") as store:
        with pytest.raises(OnyxError, match="Unable to load") as error:
            connector.upload_files([upload])
        assert error.value.error_code == OnyxErrorCode.INVALID_INPUT
        store.return_value.save_file.assert_not_called()


def test_project_upload_uses_archive_expansion_cap(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(connector, "MAX_UNZIPPED_BYTES", 8)
    upload: UploadFile = UploadFile(
        file=_archive({"a.txt": b"12345", "b.txt": b"6789"}), filename="files.zip"
    )
    categorized: MagicMock = MagicMock(acceptable=[upload])
    with (
        patch.object(projects, "categorize_uploaded_files", return_value=categorized),
        patch.object(connector, "get_default_file_store") as store,
    ):
        store.return_value.save_file.side_effect = AssertionError(
            "archive contents reached storage"
        )
        with pytest.raises(OnyxError, match="too large"):
            projects.create_user_files([upload], None, User(id=uuid4()), MagicMock())
        store.return_value.save_file.assert_not_called()


@pytest.mark.parametrize("metadata_member", [False, True])
def test_connector_limits_actual_aggregate_bytes(
    monkeypatch: pytest.MonkeyPatch, metadata_member: bool
) -> None:
    monkeypatch.setattr(connector, "MAX_UNZIPPED_BYTES", 6)
    name: str = ONYX_METADATA_FILENAME if metadata_member else "a.txt"
    upload: UploadFile = UploadFile(
        file=_archive({name: b"a", "b.txt": b"b"}), filename="files.zip"
    )
    with (
        patch.object(
            zipfile.ZipFile, "open", side_effect=[BytesIO(b"1234"), BytesIO(b"5678")]
        ),
        patch.object(connector, "get_default_file_store") as store,
    ):
        store.return_value.save_file.return_value = "stored-file"
        with pytest.raises(OnyxError, match="too large"):
            connector.upload_files([upload])
        store.return_value.save_file.assert_called_once()


def test_library_skips_oversized_member_before_decompression(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(library, "USER_LIBRARY_MAX_FILE_SIZE_BYTES", 4)
    upload: UploadFile = UploadFile(
        file=_archive({"huge.txt": b"12345", "small.txt": b"1234"}),
        filename="files.zip",
    )
    original_open = zipfile.ZipFile.open

    def open_member(
        archive: zipfile.ZipFile,
        name: str | zipfile.ZipInfo,
        mode: Literal["r", "w"] = "r",
        pwd: bytes | None = None,
        *,
        force_zip64: bool = False,
    ) -> IO[bytes]:
        member_name: str = name.filename if isinstance(name, zipfile.ZipInfo) else name
        assert member_name != "huge.txt", "oversized member was decompressed"
        return original_open(archive, name, mode, pwd, force_zip64=force_zip64)

    with (
        patch.object(library, "get_user_storage_bytes", return_value=0),
        patch.object(library, "get_or_create_craft_connector", return_value=(1, 2)),
        patch.object(
            library, "store_user_file", return_value=("doc", None, None)
        ) as store,
        patch.object(library, "create_directory_record"),
        patch.object(library, "update_connector_credential_pair"),
        patch.object(library, "cleanup_old_blobs"),
        patch.object(library, "sync_user_library_to_active_sandboxes"),
        patch.object(zipfile.ZipFile, "open", open_member),
    ):
        response = asyncio.run(
            library.upload_zip(upload, "/", User(id=uuid4()), MagicMock())
        )
    assert response.total_uploaded == 1
    assert response.entries[0].name == "small.txt"
    assert store.call_args.kwargs["content"] == b"1234"


@pytest.mark.parametrize(
    "frontmatter", ["a: &a [1]\nb: *a", "a: &a {b: 1}\nc: {<<: *a}"]
)
def test_skill_rejects_yaml_aliases(frontmatter: str) -> None:
    with pytest.raises(OnyxError, match="aliases") as error:
        metadata.parse_skill_md_frontmatter(f"---\n{frontmatter}\n---\nBody".encode())
    assert error.value.error_code == OnyxErrorCode.INVALID_INPUT


def test_skill_frontmatter_cap_preserves_markdown(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(metadata, "_MAX_FRONTMATTER_CHARS", 8)
    assert metadata.parse_skill_md_frontmatter(b"---\na: 12345\n---\n" + b"x" * 32) == (
        {"a": 12345},
        "x" * 32,
    )
    with pytest.raises(OnyxError, match="too large"):
        metadata.parse_skill_md_frontmatter(b"---\na: 123456\n---\nBody")
